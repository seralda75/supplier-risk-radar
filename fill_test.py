"""Rellena el Excel de test con las respuestas del sistema (el fichero que puntúan los moderadores).

    python fill_test.py            # con el juez LLM
    python fill_test.py --no-llm   # solo reglas

Solo se escriben las celdas de Is_Risk, Urgent y Affected_Supplier_IDs, sobre una copia del
original: la hoja README, las fechas y el formato se conservan. Después se relee el fichero y se
valida el formato exacto de la entrega. El detalle (explicación, acción, evidencia) va aparte.
"""
from __future__ import annotations

import argparse
import re
import sys

import openpyxl
import pandas as pd

from core.config import LLM_MODEL, OUTPUT_DIR, TEST_OUTPUT_XLSX, TEST_XLSX
from core.data import load_news
from core.pipeline import assess
from core.suppliers import load_suppliers

SHEET = "News_Test"
ANSWER_COLS = ["Is_Risk", "Urgent", "Affected_Supplier_IDs"]
IDS_RE = re.compile(r"^S\d{3}(; S\d{3})*$")


def write_answers(answers: dict[str, dict]) -> None:
    wb = openpyxl.load_workbook(TEST_XLSX)
    ws = wb[SHEET]
    header = {c.value: c.column for c in ws[1] if c.value}
    missing = [c for c in ["News_ID", *ANSWER_COLS] if c not in header]
    if missing:
        sys.exit(f"Faltan columnas en {SHEET}: {missing}")
    for row in range(2, ws.max_row + 1):
        nid = ws.cell(row, header["News_ID"]).value
        if nid in answers:
            for col in ANSWER_COLS:
                ws.cell(row, header[col]).value = answers[nid][col] or None   # IDs vacíos -> celda vacía
    wb.save(TEST_OUTPUT_XLSX)


def validate_output(expected_ids: list[str]) -> list[str]:
    """Relee el fichero generado y devuelve la lista de problemas de formato (vacía si todo está bien)."""
    df = pd.read_excel(TEST_OUTPUT_XLSX, sheet_name=SHEET, dtype=str, keep_default_na=False)
    original = pd.read_excel(TEST_XLSX, sheet_name=SHEET, dtype=str, keep_default_na=False)
    known = set(load_suppliers())
    problems = []
    if df.News_ID.tolist() != expected_ids:
        problems.append("Las noticias no coinciden con el original")
    keep = [c for c in original.columns if c not in ANSWER_COLS]
    if not df[keep].equals(original[keep]):
        problems.append("Se ha modificado alguna columna que no es de respuesta")
    for r in df.itertuples():
        if r.Is_Risk not in ("Y", "N") or r.Urgent not in ("Y", "N"):
            problems.append(f"{r.News_ID}: Is_Risk/Urgent deben ser Y o N")
        if r.Is_Risk == "N" and r.Urgent == "Y":
            problems.append(f"{r.News_ID}: solo un riesgo puede ser urgente")
        ids = r.Affected_Supplier_IDs
        if ids and (not IDS_RE.match(ids) or not set(ids.split("; ")) <= known):
            problems.append(f"{r.News_ID}: IDs con formato incorrecto o inexistentes: '{ids}'")
    wb = openpyxl.load_workbook(TEST_OUTPUT_XLSX)
    if wb.sheetnames != openpyxl.load_workbook(TEST_XLSX).sheetnames:
        problems.append("Las hojas del fichero no coinciden con el original")
    return problems


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--no-llm", action="store_true", help="solo reglas (sin Ollama)")
    ap.add_argument("--model", default=LLM_MODEL)
    args = ap.parse_args()
    use_llm = not args.no_llm
    if use_llm:
        from core import llm
        if not llm.is_up():
            sys.exit("Ollama no responde. Arráncalo (ollama serve) o usa --no-llm.")

    news = load_news("test")
    answers, detail = {}, []
    for n in news:
        r = assess(n, use_llm=use_llm, model=args.model, log=True)
        answers[n["News_ID"]] = r.to_row()
        detail.append({"News_ID": r.news_id, "Headline": r.headline, **r.to_row(), "Event_Type": r.event_label,
                       "Severity": r.severity, "Recovery_Days": r.recovery_days, "Confidence": r.confidence,
                       "Needs_Review": r.needs_review, "Explanation": r.explanation, "Action": r.action,
                       "Evidence": r.evidence, "Mode": r.mode})
        print(f"  {r.news_id} risk={answers[r.news_id]['Is_Risk']} urg={answers[r.news_id]['Urgent']} "
              f"ids=[{answers[r.news_id]['Affected_Supplier_IDs']}] {r.event_type}", flush=True)

    write_answers(answers)
    detail_path = OUTPUT_DIR / "test_detalle.csv"
    pd.DataFrame(detail).to_csv(detail_path, index=False, encoding="utf-8-sig")
    problems = validate_output([n["News_ID"] for n in news])
    d = pd.DataFrame(detail)
    print(f"\nRiesgos: {(d.Is_Risk == 'Y').sum()}/{len(d)} · urgentes: {(d.Urgent == 'Y').sum()}")
    print(f"Entrega: {TEST_OUTPUT_XLSX}\nDetalle: {detail_path}")
    if problems:
        print("\n❌ Problemas de formato:\n  " + "\n  ".join(problems))
        sys.exit(1)
    print("✅ Formato de la entrega validado")


if __name__ == "__main__":
    main()
